"""Validated private uploads and bounded model-readable document extraction."""

import base64
import hashlib
import io
import json
import warnings
import zipfile
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import connection, transaction
from django.db.models import F

from defusedxml import ElementTree
from PIL import Image, ImageOps

from .assistant_attachment_models import AssistantChatAttachment
from .assistant_gateway import DeliveryConflictError

FILE_LIMIT = 20 * 1024 * 1024
MESSAGE_FILE_LIMIT = 40 * 1024 * 1024
ATTACHMENT_LIMIT = 8
TEXT_LIMIT = 120000
IMAGE_PIXEL_LIMIT = 20_000_000
EXTENSIONS = {
    '.txt': 'text/plain',
    '.md': 'text/markdown',
    '.log': 'text/plain',
    '.csv': 'text/csv',
    '.tsv': 'text/tab-separated-values',
    '.json': 'application/json',
    '.pdf': 'application/pdf',
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
}
IMAGE_TYPES = {
    'JPEG': 'image/jpeg',
    'PNG': 'image/png',
    'WEBP': 'image/webp',
    'GIF': 'image/gif',
}


def attachment_metadata(item):
    """Only authenticated application endpoints are exposed to the browser.

    URLs carry the owning partner selector so previews and draft restoration
    keep working after switching partners; selector-less calls still resolve
    the default partner for backward compatibility.
    """
    route = f'/api/cangxu/assistant/attachments/{item.pk}'
    selector = f'?partner_id={item.binding.partner_id}'
    return {
        'id': str(item.pk),
        'name': item.name,
        'mime_type': item.mime_type,
        'size': item.size,
        'kind': item.kind,
        'image_width': item.image_width,
        'image_height': item.image_height,
        'download_url': f'{route}/download/{selector}',
        'preview_url': f'{route}/preview/{selector}' if item.kind == 'image' else None,
    }


def _validated_image(data):
    """Decode images rather than trusting extension or browser MIME headers."""
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(data)) as image:
            if (
                image.format not in IMAGE_TYPES
                or image.width * image.height > IMAGE_PIXEL_LIMIT
            ):
                raise ValidationError('Image format or dimensions are unsupported')
            mime, width, height = IMAGE_TYPES[image.format], image.width, image.height
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            image.load()
    return mime, width, height


def _safe_archive(data, required):
    """Bound package expansion before parsing Office XML."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        files = archive.infolist()
        if (
            len(files) > 2000
            or sum(item.file_size for item in files) > 40 * 1024 * 1024
        ):
            raise ValidationError('Office document expands beyond the supported limit')
        if required not in archive.namelist() or any(
            item.flag_bits & 1 for item in files
        ):
            raise ValidationError('Unsupported or encrypted Office document')
        for item in files:
            if item.file_size > 0 and item.compress_size == 0:
                raise ValidationError('Invalid Office package')
        return archive.read(required)


def create_upload(binding, upload, client_upload_id):
    """Validate contents and persist exactly one file for each upload key."""
    if upload.size <= 0 or upload.size > FILE_LIMIT:
        raise ValidationError('Provide a nonempty file up to 20 MiB')
    data = upload.read(FILE_LIMIT + 1)
    if len(data) > FILE_LIMIT:
        raise ValidationError('File exceeds 20 MiB')
    name = Path(upload.name.replace('\\', '/')).name[:255]
    extension = Path(name).suffix.lower()
    kind, width, height = 'document', None, None
    try:
        if extension in {'.png', '.jpg', '.jpeg', '.gif', '.webp'}:
            mime, width, height = _validated_image(data)
            kind = 'image'
        elif extension in EXTENSIONS:
            mime = EXTENSIONS[extension]
            if extension == '.pdf':
                from pypdf import PdfReader

                document = PdfReader(io.BytesIO(data))
                if document.is_encrypted or len(document.pages) > 300:
                    raise ValidationError(
                        'Encrypted PDFs or PDFs over 300 pages are unsupported'
                    )
            elif extension == '.docx':
                ElementTree.fromstring(_safe_archive(data, 'word/document.xml'))
            elif extension == '.xlsx':
                ElementTree.fromstring(_safe_archive(data, 'xl/workbook.xml'))
            else:
                text = data.decode('utf-8-sig')
                if '\x00' in text:
                    raise ValidationError('Text files must contain UTF-8 text')
                if extension == '.json':
                    json.loads(text)
        else:
            raise ValidationError(
                'Supported: PNG/JPEG/WEBP/GIF, PDF/DOCX/XLSX, UTF-8 TXT/MD/CSV/TSV/JSON/LOG'
            )
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(
            'File content does not match a supported readable format'
        ) from exc
    digest = hashlib.sha256(data).hexdigest()
    from .assistant_database import sqlite_busy_retry

    return sqlite_busy_retry(
        lambda: _persist_upload(
            binding, client_upload_id, data, name, mime, kind, width, height, digest
        )
    )


def _persist_upload(
    binding, client_upload_id, data, name, mime, kind, width, height, digest
):
    """Each retry owns a fresh transaction and cleans any uncommitted file copy."""
    item = None
    try:
        with transaction.atomic():
            # The same owner lock also serializes concurrent uploads and message assembly.
            if connection.vendor == 'sqlite':
                # SQLite has no row locks: obtain its writer reservation before reading.
                # This avoids the impossible deferred read-to-write lock upgrade race.
                type(binding).objects.filter(pk=binding.pk).update(
                    updated_at=F('updated_at')
                )
            type(binding).objects.select_for_update().get(pk=binding.pk)
            existing = AssistantChatAttachment.objects.filter(
                binding=binding, client_upload_id=client_upload_id
            ).first()
            if existing:
                if (existing.sha256, existing.name) != (digest, name):
                    raise DeliveryConflictError(
                        'Upload key already identifies different content'
                    )
                return existing, False
            item = AssistantChatAttachment(
                binding=binding,
                client_upload_id=client_upload_id,
                name=name,
                mime_type=mime,
                size=len(data),
                sha256=digest,
                kind=kind,
                image_width=width,
                image_height=height,
            )
            item.file.save(str(item.id), ContentFile(data), save=False)
            item.save()
            return item, True
    except Exception:
        if item is not None and item.file.name:
            item.file.storage.delete(item.file.name)
        raise


def message_attachments(binding, identifiers):
    """Resolve each upload against the current server-owned partner."""
    if len(identifiers) > ATTACHMENT_LIMIT or len(set(identifiers)) != len(identifiers):
        raise ValidationError('Provide up to eight distinct attachments')
    found = {
        str(item.pk): item
        for item in AssistantChatAttachment.objects.filter(
            binding=binding, pk__in=identifiers
        ).select_related('binding')
    }
    if len(found) != len(identifiers):
        raise ValidationError('One or more chat attachments are unavailable')
    items = [found[str(identifier)] for identifier in identifiers]
    if sum(item.size for item in items) > MESSAGE_FILE_LIMIT:
        raise ValidationError('Combined message attachments exceed 40 MiB')
    return [attachment_metadata(item) for item in items]


def model_attachment_content(item):
    """Return genuine bounded text or image data, with explicit extraction limits."""
    with item.file.open('rb') as source:
        data = source.read(FILE_LIMIT + 1)
    if len(data) != item.size or hashlib.sha256(data).hexdigest() != item.sha256:
        raise ValidationError('Stored attachment failed integrity verification')
    result = {**attachment_metadata(item), 'content': [], 'extraction_status': 'parsed'}
    if item.kind == 'image':
        _validated_image(data)
        with Image.open(io.BytesIO(data)) as image:
            image.seek(0)
            image = ImageOps.exif_transpose(image)
            image.thumbnail((2048, 2048))
            canvas = Image.new('RGB', image.size, 'white')
            rgba = image.convert('RGBA')
            canvas.paste(rgba, mask=rgba.getchannel('A'))
            output = io.BytesIO()
            canvas.save(output, format='JPEG', quality=90)
        result['content'] = [
            {
                'type': 'image',
                'data': base64.b64encode(output.getvalue()).decode('ascii'),
                'mimeType': 'image/jpeg',
            }
        ]
        result['extraction_note'] = (
            'First frame only; scaled to at most 2048 pixels per dimension for model input. Download retains original.'
        )
        return result
    extension = Path(item.name).suffix.lower()
    try:
        if extension == '.pdf':
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pieces, length, has_pdf_text = [], 0, False
            for number, page in enumerate(reader.pages):
                extracted = page.extract_text() or ''
                has_pdf_text = has_pdf_text or bool(extracted.strip())
                piece = f'\n[Page {number + 1}]\n{extracted}'
                pieces.append(piece[: TEXT_LIMIT - length])
                length += len(piece)
                if length >= TEXT_LIMIT:
                    break
            text = ''.join(pieces)
            if not has_pdf_text:
                result['extraction_status'] = 'unsupported'
                result['extraction_note'] = (
                    'PDF has no extractable text; scanned PDF OCR is not available.'
                )
        elif extension == '.docx':
            xml = _safe_archive(data, 'word/document.xml')
            tree = ElementTree.fromstring(xml)
            text = '\n'.join(
                node.text or '' for node in tree.iter() if node.tag.endswith('}t')
            )
        elif extension == '.xlsx':
            import openpyxl

            _safe_archive(data, 'xl/workbook.xml')
            book = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=False
            )
            lines, length = [], 0
            try:
                for sheet in book:
                    lines.append(f'\n[Sheet: {sheet.title}]')
                    if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 100:
                        result['extraction_status'] = 'truncated'
                        result['extraction_note'] = (
                            'Spreadsheet input includes at most 10000 rows and 100 columns per sheet; download the original for all cells.'
                        )
                    for row in sheet.iter_rows(
                        max_row=min(sheet.max_row or 1, 10000),
                        max_col=min(sheet.max_column or 1, 100),
                        values_only=True,
                    ):
                        line = '\t'.join(
                            str(value) if value is not None else ''
                            for value in row[:100]
                        )
                        lines.append(line[: TEXT_LIMIT - length])
                        length += len(line) + 1
                        if length >= TEXT_LIMIT:
                            break
                    if length >= TEXT_LIMIT:
                        break
                text = '\n'.join(lines)
            finally:
                book.close()
        else:
            text = data.decode('utf-8-sig')
        if len(text) >= TEXT_LIMIT:
            result['extraction_status'] = 'truncated'
            result['extraction_note'] = (
                'Only the first 120000 characters are included; use the original file for complete content.'
            )
        result['content'] = [{'type': 'text', 'text': text[:TEXT_LIMIT]}]
    except Exception:
        result['extraction_status'] = 'failed'
        result['extraction_note'] = (
            'Document extraction failed; no document content was provided to the model.'
        )
    return result
