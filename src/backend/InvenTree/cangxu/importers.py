"""Read-only workbook parsing helpers for staged Cangxu imports."""

from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_WORKBOOK_CELLS = 500_000
MAX_WORKBOOK_ROWS = 50_000


class WorkbookImportError(ValueError):
    """Raised when an uploaded workbook cannot be safely staged."""


def _json_value(value):
    """Convert supported Excel values into stable JSON-compatible scalars."""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def parse_workbook_rows(upload):
    """Return populated worksheet rows while preserving cell coordinates/format."""
    if upload.size > MAX_WORKBOOK_BYTES:
        raise WorkbookImportError('workbook_too_large')

    try:
        workbook = load_workbook(upload, read_only=True, data_only=False)
    except Exception as exc:
        raise WorkbookImportError('invalid_xlsx') from exc

    rows = []
    cell_count = 0
    try:
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(), start=1):
                cell_count += len(row)
                if cell_count > MAX_WORKBOOK_CELLS:
                    raise WorkbookImportError('workbook_too_many_cells')

                populated = [
                    (column_index, cell)
                    for column_index, cell in enumerate(row, start=1)
                    if cell.value is not None
                ]
                if not populated:
                    continue

                values = {
                    get_column_letter(column_index): {
                        'value': _json_value(cell.value),
                        'data_type': cell.data_type,
                        'number_format': cell.number_format,
                    }
                    for column_index, cell in enumerate(row, start=1)
                }
                first_col = min(column for column, _cell in populated)
                last_col = max(column for column, _cell in populated)
                rows.append({
                    'source_sheet': sheet.title[:100],
                    'source_row': row_index,
                    'source_range': (
                        f'{get_column_letter(first_col)}{row_index}:'
                        f'{get_column_letter(last_col)}{row_index}'
                    ),
                    'source_values': values,
                    'normalized_values': {},
                })
                if len(rows) > MAX_WORKBOOK_ROWS:
                    raise WorkbookImportError('workbook_too_many_rows')
    except WorkbookImportError:
        raise
    except Exception as exc:
        raise WorkbookImportError('invalid_xlsx') from exc
    finally:
        workbook.close()

    if not rows:
        raise WorkbookImportError('workbook_has_no_populated_rows')

    return rows
